import re
import xml.etree.ElementTree as ET
from collections import OrderedDict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from zipfile import ZipFile

from task_data_model import Step, TaskPlanData


REQUIRED_COLUMNS = {
    "步骤序号": "step_index",
    "设备ID(Hex)": "device_id_hex",
    "指令ID(Hex)": "cmd_id_hex",
    "指令名称": "cmd_name",
    "参数排序": "param_order",
    "参数名称": "param_name",
    "数据类型": "data_type",
    "赋值": "value",
}


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_ref_to_index(cell_ref: str) -> int:
    letters = re.sub(r"[^A-Z]", "", cell_ref.upper())
    index = 0
    for ch in letters:
        index = index * 26 + (ord(ch) - ord("A") + 1)
    return index - 1


def _read_xlsx_with_zip(file_path: str, sheet_name: Optional[str] = None) -> List[List[str]]:
    ns = {
        "a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    path = Path(file_path)

    with ZipFile(path) as archive:
        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for si in root.findall("a:si", ns):
                parts = []
                for t in si.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"):
                    parts.append(t.text or "")
                shared_strings.append("".join(parts))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.get("Id"): rel.get("Target")
            for rel in rels.findall("rel:Relationship", ns)
        }

        selected_target = None
        available_sheets = []
        for sheet in workbook.findall(".//a:sheets/a:sheet", ns):
            name = sheet.get("name", "")
            available_sheets.append(name)
            rel_id = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = rel_targets.get(rel_id)
            if selected_target is None and sheet_name is None:
                selected_target = target
            if sheet_name is not None and name == sheet_name:
                selected_target = target

        if selected_target is None:
            raise ValueError(f"worksheet '{sheet_name}' not found; available: {available_sheets}")

        selected_target = selected_target.lstrip("/")
        if not selected_target.startswith("xl/"):
            selected_target = f"xl/{selected_target}"

        if selected_target not in archive.namelist():
            raise ValueError(f"worksheet XML not found: {selected_target}")

        root = ET.fromstring(archive.read(selected_target))
        rows: List[List[str]] = []

        for row in root.findall(".//a:sheetData/a:row", ns):
            values: Dict[int, str] = {}
            max_col = -1

            for cell in row.findall("a:c", ns):
                cell_ref = cell.get("r", "")
                col_index = _cell_ref_to_index(cell_ref)
                max_col = max(max_col, col_index)
                cell_type = cell.get("t", "")
                text = ""

                if cell_type == "s":
                    value = cell.find("a:v", ns)
                    if value is not None and value.text is not None:
                        text = shared_strings[int(value.text)]
                elif cell_type == "inlineStr":
                    texts = []
                    for t in cell.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"):
                        texts.append(t.text or "")
                    text = "".join(texts)
                else:
                    value = cell.find("a:v", ns)
                    text = "" if value is None or value.text is None else value.text

                values[col_index] = text

            if max_col >= 0:
                rows.append([values.get(i, "") for i in range(max_col + 1)])

        return rows


def _read_xlsx_rows(file_path: str, sheet_name: Optional[str] = None) -> List[List[str]]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return _read_xlsx_with_zip(file_path, sheet_name=sheet_name)

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"worksheet '{sheet_name}' not found; available: {workbook.sheetnames}")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.worksheets[0]

        return [[_cell_to_text(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _normalize_header(header: str) -> str:
    return re.sub(r"\s+", "", header.strip())


def _parse_int(value: Any, name: str) -> int:
    text = _cell_to_text(value)
    if text == "":
        raise ValueError(f"{name} is empty")
    return int(text, 0)


def _parse_hex_int(value: Any, name: str) -> int:
    text = _cell_to_text(value)
    if text == "":
        raise ValueError(f"{name} is empty")
    return int(text, 16) if not text.lower().startswith("0x") else int(text, 16)


def _clean_hex(text: str) -> str:
    if text.lower().startswith("0x"):
        text = text[2:]
    return re.sub(r"[\s:_,-]", "", text)


def _convert_value_to_bytes(value: Any, data_type: str, row_num: int) -> bytes:
    dtype = _cell_to_text(data_type).lower().replace(" ", "")
    text = _cell_to_text(value)

    if dtype == "" and text == "":
        return b""

    if dtype == "raw_hex":
        hex_text = _clean_hex(text)
        if len(hex_text) == 0:
            return b""
        if len(hex_text) % 2 != 0:
            raise ValueError(f"row {row_num}: raw_hex must contain an even number of digits")
        return bytes.fromhex(hex_text)

    if dtype in ("uint8", "uint8_t", "byte"):
        val = _parse_int(value, f"row {row_num} uint8")
        if not 0 <= val <= 0xFF:
            raise ValueError(f"row {row_num}: uint8 out of range: {val}")
        return val.to_bytes(1, "big")

    if dtype in ("uint16", "uint16_t", "word"):
        val = _parse_int(value, f"row {row_num} uint16")
        if not 0 <= val <= 0xFFFF:
            raise ValueError(f"row {row_num}: uint16 out of range: {val}")
        return val.to_bytes(2, "big")

    if dtype in ("int16", "int16_t"):
        val = _parse_int(value, f"row {row_num} int16")
        if not -0x8000 <= val <= 0x7FFF:
            raise ValueError(f"row {row_num}: int16 out of range: {val}")
        return val.to_bytes(2, "big", signed=True)

    if dtype in ("uint32", "uint32_t", "dword"):
        val = _parse_int(value, f"row {row_num} uint32")
        if not 0 <= val <= 0xFFFFFFFF:
            raise ValueError(f"row {row_num}: uint32 out of range: {val}")
        return val.to_bytes(4, "big")

    if dtype in ("int32", "int32_t"):
        val = _parse_int(value, f"row {row_num} int32")
        if not -0x80000000 <= val <= 0x7FFFFFFF:
            raise ValueError(f"row {row_num}: int32 out of range: {val}")
        return val.to_bytes(4, "big", signed=True)

    raise ValueError(f"row {row_num}: unsupported data type '{data_type}'")


def parse_commands_to_list(file_path: str = "commands.xlsx",
                           sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    rows = _read_xlsx_rows(file_path, sheet_name)
    rows = [row for row in rows if any(_cell_to_text(cell) for cell in row)]

    if not rows:
        return []

    headers = [_normalize_header(_cell_to_text(header)) for header in rows[0]]
    column_mapping: Dict[str, int] = {}

    for expected, key in REQUIRED_COLUMNS.items():
        expected_norm = _normalize_header(expected)
        if expected_norm not in headers:
            raise ValueError(f"missing required column: {expected}")
        column_mapping[key] = headers.index(expected_norm)

    result: List[Dict[str, Any]] = []
    for row_num, row in enumerate(rows[1:], start=2):
        item: Dict[str, Any] = {"row_num": row_num}
        for key, idx in column_mapping.items():
            item[key] = _cell_to_text(row[idx]) if idx < len(row) else ""
        result.append(item)

    return result


def convert_commands_to_steps(commands_list: Iterable[Dict[str, Any]]) -> List[Step]:
    grouped: "OrderedDict[int, List[Dict[str, Any]]]" = OrderedDict()

    for row in commands_list:
        step_index = _parse_int(row.get("step_index"), f"row {row.get('row_num')} step_index")
        if step_index <= 0:
            raise ValueError(f"row {row.get('row_num')}: step index must be positive")
        grouped.setdefault(step_index, []).append(row)

    steps: List[Step] = []

    for step_index in sorted(grouped.keys()):
        rows = grouped[step_index]
        first = rows[0]
        device_id = _parse_hex_int(first["device_id_hex"], f"step {step_index} device_id")
        cmd_id = _parse_hex_int(first["cmd_id_hex"], f"step {step_index} cmd_id")

        if not 0 <= device_id <= 0xFF:
            raise ValueError(f"step {step_index}: device_id out of range: {device_id}")
        if not 0 <= cmd_id <= 0xFFFF:
            raise ValueError(f"step {step_index}: cmd_id out of range: {cmd_id}")

        args_parts = []
        seen_orders = set()

        for row in rows:
            row_num = row.get("row_num")
            row_device_id = _parse_hex_int(row["device_id_hex"], f"row {row_num} device_id")
            row_cmd_id = _parse_hex_int(row["cmd_id_hex"], f"row {row_num} cmd_id")
            if row_device_id != device_id or row_cmd_id != cmd_id:
                raise ValueError(f"step {step_index}: mixed device_id/cmd_id in row {row_num}")

            data_type = _cell_to_text(row.get("data_type"))
            value = _cell_to_text(row.get("value"))
            param_order_text = _cell_to_text(row.get("param_order"))

            if data_type == "" and value == "" and param_order_text == "":
                continue

            param_order = _parse_int(param_order_text, f"row {row_num} param_order")
            if param_order in seen_orders:
                raise ValueError(f"step {step_index}: duplicate param_order {param_order}")
            seen_orders.add(param_order)

            args_parts.append((param_order, _convert_value_to_bytes(value, data_type, row_num)))

        args = b"".join(part for _, part in sorted(args_parts, key=lambda item: item[0]))
        steps.append(Step(device_id=device_id, cmd_id=cmd_id, args=args))

    return steps


def parse_excel_to_task_plan(file_path: str = "commands.xlsx",
                             sheet_name: Optional[str] = None,
                             format_ver: int = 1,
                             plan_id: int = 100) -> TaskPlanData:
    commands = parse_commands_to_list(file_path, sheet_name)
    steps = convert_commands_to_steps(commands)
    return TaskPlanData(format_ver=format_ver, plan_id=plan_id, steps=steps)


if __name__ == "__main__":
    plan = parse_excel_to_task_plan("commands.xlsx", plan_id=100)
    data = plan.to_bytes()
    print(plan)
    print(" ".join(f"{byte:02X}" for byte in data))
